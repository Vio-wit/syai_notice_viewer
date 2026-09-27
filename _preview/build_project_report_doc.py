from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "PROJECT_REPORT.md"
OUTPUT = ROOT / "PROJECT_REPORT.docx"
SCREENSHOTS = [
    (ROOT / "output" / "playwright" / "dashboard.png", "图 1：本地看板的比赛关键词筛选结果"),
    (ROOT / "output" / "playwright" / "notice-detail.png", "图 2：点击查看详情后的通知弹窗"),
]


def set_cell_text(cell, text: str) -> None:
    cell.text = text.strip().replace("`", "")
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Microsoft YaHei"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            run.font.size = Pt(9)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.append(field_begin)
    run._r.append(instruction)
    run._r.append(field_end)
    paragraph.add_run(" 页")


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color in (("Title", 24, "17365D"), ("Heading 1", 16, "1F4E79"), ("Heading 2", 13, "2F75B5"), ("Heading 3", 11, "5B9BD5")):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)

    if "CaptionCN" not in styles:
        caption = styles.add_style("CaptionCN", WD_STYLE_TYPE.PARAGRAPH)
        caption.font.name = "Microsoft YaHei"
        caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        caption.font.size = Pt(9)
        caption.font.color.rgb = RGBColor(96, 96, 96)
    else:
        caption = styles["CaptionCN"]
    caption.paragraph_format.space_after = Pt(10)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    for run in footer.runs:
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(128, 128, 128)


def add_cover(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(90)
    run = paragraph.add_run("学院通知查看器")
    run.bold = True
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(23, 54, 93)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("项目建设与运行报告")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(47, 117, 181)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(24)
    paragraph.add_run("报告快照：2026-09-28\n").font.size = Pt(11)
    paragraph.add_run("项目目录：K:\\syai_notice_viewer\n").font.size = Pt(11)
    paragraph.add_run("本地看板：http://127.0.0.1:8765/").font.size = Pt(11)
    document.add_page_break()


def add_screenshots(document: Document) -> None:
    document.add_heading("看板真实截图", level=1)
    document.add_paragraph("以下图片由本地运行中的看板实际生成，展示关键词筛选和通知详情弹窗。截图仅用于说明当前界面，不代表额外的演示数据。")
    for image_path, caption in SCREENSHOTS:
        if not image_path.exists():
            continue
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(image_path), width=Inches(6.25))
        caption_paragraph = document.add_paragraph(caption, style="CaptionCN")
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_page_break()


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = document.add_table(rows=1, cols=len(rows[0]))
    table.style = "Light Shading Accent 1"
    for index, value in enumerate(rows[0]):
        set_cell_text(table.rows[0].cells[index], value)
    for row in rows[1:]:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
    document.add_paragraph()


def clean_inline(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)
    text = text.replace("`", "")
    return text.strip()


def add_report_markdown(document: Document) -> None:
    lines = REPORT.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if line.startswith("# "):
            index += 1
            continue
        if line.startswith("## ") or line.startswith("### "):
            level = 1 if line.startswith("## ") else 2
            document.add_heading(clean_inline(line.split(" ", 1)[1]), level=level)
            index += 1
            continue
        if line.startswith("```"):
            code_lines: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                code_lines.append(lines[index])
                index += 1
            paragraph = document.add_paragraph()
            paragraph.style = "No Spacing"
            run = paragraph.add_run("\n".join(code_lines))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            index += 1
            continue
        if line.startswith("|"):
            table_lines: list[str] = []
            while index < len(lines) and lines[index].startswith("|"):
                current = lines[index]
                if not re.match(r"^\|\s*:?-{2,}", current):
                    table_lines.append(current)
                index += 1
            rows = [[clean_inline(cell) for cell in row.strip().strip("|").split("|")] for row in table_lines]
            add_table(document, rows)
            continue
        if line.startswith("- "):
            paragraph = document.add_paragraph(clean_inline(line[2:]), style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(2)
            index += 1
            continue
        if re.match(r"^\d+\. ", line):
            paragraph = document.add_paragraph(clean_inline(re.sub(r"^\d+\. ", "", line)), style="List Number")
            paragraph.paragraph_format.space_after = Pt(2)
            index += 1
            continue
        if line.strip():
            document.add_paragraph(clean_inline(line))
        index += 1


def main() -> None:
    document = Document()
    configure_document(document)
    document.core_properties.title = "学院通知查看器项目报告"
    document.core_properties.subject = "项目建设与运行报告"
    document.core_properties.author = "Codex"
    add_cover(document)
    add_screenshots(document)
    add_report_markdown(document)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
